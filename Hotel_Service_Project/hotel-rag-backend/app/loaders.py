"""Document loaders for processing .docx files."""

import os
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from docx import Document
from langchain_core.documents import Document as LangChainDocument

from app.config import settings

# Category mapping from doc_collection to category
CATEGORY_MAPPING = {
    "corporate": "policies",
    "luxury-suites": "rooms_amenities",
    "waypoint-inns": "local_guide",
    "family-getaways": "services",
    "party-times": "services",
    "seaside-resorts": "services",
}


def get_category(doc_collection: str) -> str:
    """Map doc_collection to category."""
    return CATEGORY_MAPPING.get(doc_collection, "general")


def get_file_modified_date(file_path: Path) -> str:
    """Get file modified date in YYYY-MM-DD format."""
    mtime = os.path.getmtime(file_path)
    return datetime.fromtimestamp(mtime).strftime("%Y-%m-%d")


def is_heading_style(style_name: Optional[str]) -> bool:
    """Check if a style name is a heading style."""
    if not style_name:
        return False
    return style_name.startswith("Heading")


def load_docx_file(file_path: Path) -> Document:
    """Load a .docx file using python-docx."""
    return Document(str(file_path))


def extract_sections(
    doc: Document, source_file: str, doc_collection: str, file_path: Path
) -> List[LangChainDocument]:
    """
    Extract sections from a Word document based on heading styles.

    Sections are defined by headings (Heading 1, Heading 2, etc.).
    All paragraphs following a heading until the next heading belong to that section.
    """
    sections = []
    current_section_title = "Untitled Section"
    current_paragraphs = []
    file_date = get_file_modified_date(file_path)

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else None

        if is_heading_style(style_name):
            # Save previous section if it has content
            if current_paragraphs:
                section_text = "\n".join(current_paragraphs).strip()
                if section_text:
                    sections.append(
                        LangChainDocument(
                            page_content=section_text,
                            metadata={
                                "audience": settings.default_audience,
                                "category": get_category(doc_collection),
                                "doc_collection": doc_collection,
                                "source_file": source_file,
                                "section_title": current_section_title,
                                "last_updated": file_date,
                            },
                        )
                    )

            # Start new section
            current_section_title = para.text.strip() or "Untitled Section"
            current_paragraphs = []
        else:
            # Add paragraph to current section
            para_text = para.text.strip()
            if para_text:
                current_paragraphs.append(para_text)

    # Don't forget the last section
    if current_paragraphs:
        section_text = "\n".join(current_paragraphs).strip()
        if section_text:
            sections.append(
                LangChainDocument(
                    page_content=section_text,
                    metadata={
                        "audience": settings.default_audience,
                        "category": get_category(doc_collection),
                        "doc_collection": doc_collection,
                        "source_file": source_file,
                        "section_title": current_section_title,
                        "last_updated": file_date,
                    },
                )
            )

    return sections


def load_document_sections(file_path: Path, doc_collection: str) -> List[LangChainDocument]:
    """
    Load a .docx file and extract sections.

    Args:
        file_path: Path to the .docx file
        doc_collection: The collection name (folder name)

    Returns:
        List of LangChain Document objects, one per section
    """
    doc = load_docx_file(file_path)
    source_file = file_path.name
    sections = extract_sections(doc, source_file, doc_collection, file_path)

    return sections

